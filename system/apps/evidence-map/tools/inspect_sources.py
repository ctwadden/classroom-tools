from pathlib import Path
import json, zipfile, hashlib, re, collections
from xml.etree import ElementTree as ET
from docx import Document
from pypdf import PdfReader
import openpyxl

BASE=Path('/Users/cdawg/Downloads')
OUT=Path('/private/tmp/evidence-map-intake')
OUT.mkdir(parents=True, exist_ok=True)
names=['Assessment_Guides (1).md','Assessment_Guides.md','Studio_Revision.json','Studio_Projects_and_Next_Five_Classes.md','Studio_Advertising_Release.zip','Evidence_Map_Review.html','START_HERE(1).html','TeachOS_Course_System_v2.zip','Course_System.json','01_MASTER_SYSTEM.md','Teaching_Year_and_App_Work_Pack.zip','Research_and_Course_Architecture.docx','Semester_Assessment_and_Evidence_Map.xlsx','App_Update_Blueprint.md','Resource_Update_Prompts_for_Claude_and_Codex.md','Resource_Review_and_Operation_7_Improvements.md','Weekend_Virtual_Production_Test.md','IB_DS_Full_Year_Plan_and_Prediction.xlsx','Outcome_Evidence_Map_Update_Prompts.md','Daily_Lessons_and_Planned_Gradebooks.xlsx','IB_Digital_Society_Day_by_Day.md','Communication_11_Day_by_Day.md','Multimedia_12_Day_by_Day.md','Weekly_Fabrication_Prompts.md','Install_This_Week_In_Your_Apps.md','Book_Resource_Audit.md','Assessment_Guides (2).md']
paths=[BASE/n for n in names]+list((BASE/'Teaching_Year_and_App_Work_Pack').rglob('*'))
paths += list(Path('/Users/cdawg/Desktop').glob('How Teens Win_*.epub'))
paths += [Path('/Users/cdawg/Desktop/Show Your Work_ Assessment in the Age of AI.pdf')]
catalog=[]; seen={}; texts={}; sheets={}
def ingest(label,data,suffix):
    sha=hashlib.sha256(data).hexdigest()
    row={'source':label,'bytes':len(data),'sha256':sha}
    if sha in seen:
        row['duplicate_of']=seen[sha];catalog.append(row);return
    seen[sha]=label;catalog.append(row)
    key=f'{len(catalog):03d}_'+re.sub(r'[^A-Za-z0-9._-]','_',Path(label).name)
    if suffix in ['.md','.json','.html','.csv']:
        text=data.decode('utf-8',errors='replace');texts[label]=text
        (OUT/key).write_text(text)
        row['extracted']=str(OUT/key)
        row['headings']=re.findall(r'^#{1,4} .*$',text,re.M)[:250]
        if suffix=='.json':
            obj=json.loads(text);row['structure']={k:len(v) if isinstance(v,(list,dict)) else str(v)[:160] for k,v in obj.items()} if isinstance(obj,dict) else {'length':len(obj)}
    elif suffix in ['.xlsx','.docx','.pdf','.epub']:
        tmp=OUT/key;tmp.write_bytes(data)
        if suffix=='.xlsx':
            wb=openpyxl.load_workbook(tmp,data_only=False,read_only=True)
            arr=[]
            for s in wb:
                rows=[list(r) for r in s.iter_rows(values_only=True)]
                arr.append({'name':s.title,'rows':s.max_row,'cols':s.max_column,'values':rows})
            sheets[label]=arr
            row['sheets']=[{k:s[k] for k in ['name','rows','cols']} for s in arr]
        elif suffix=='.docx':
            d=Document(tmp);text='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(' | '.join(c.text for c in r.cells) for t in d.tables for r in t.rows)
            texts[label]=text;(OUT/(key+'.txt')).write_text(text)
        elif suffix=='.pdf':
            d=PdfReader(tmp);text='\n'.join(f'\nPAGE {i+1}\n'+(p.extract_text() or '') for i,p in enumerate(d.pages));row['pages']=len(d.pages)
            texts[label]=text;(OUT/(key+'.txt')).write_text(text)
        else:
            z=zipfile.ZipFile(tmp);parts=[]
            for f in z.namelist():
                if f.endswith(('.xhtml','.html','.opf','.ncx')):
                    try: parts.append(f+'\n'+' '.join(ET.fromstring(z.read(f)).itertext()))
                    except ET.ParseError: pass
            text='\n'.join(parts);texts[label]=text;(OUT/(key+'.txt')).write_text(text)
for p in paths:
    if not p.exists():catalog.append({'source':str(p),'missing':True});continue
    if p.is_dir():continue
    if p.suffix=='.zip':
        ingest(str(p),p.read_bytes(),p.suffix)
        with zipfile.ZipFile(p) as z:
            for f in z.infolist():
                if f.is_dir() or f.filename.startswith('__MACOSX/') or Path(f.filename).name.startswith('.'):continue
                if f.file_size>20_000_000:continue
                ingest(str(p)+'::'+f.filename,z.read(f),Path(f.filename).suffix)
    else:ingest(str(p),p.read_bytes(),p.suffix)
(OUT/'catalog.json').write_text(json.dumps(catalog,indent=2))
(OUT/'texts.json').write_text(json.dumps(texts))
(OUT/'sheets.json').write_text(json.dumps(sheets,default=str))
print(json.dumps({'sources':len(catalog),'unique':len(seen),'duplicates':sum('duplicate_of' in x for x in catalog),'missing':[x for x in catalog if x.get('missing')],'text_files':len(texts),'workbooks':{k:[(s['name'],s['rows'],s['cols']) for s in v] for k,v in sheets.items()}},indent=2))
